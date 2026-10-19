import os
from pathlib import Path

import pytest

from tagscan.catalog import Catalog, normalize_number
from tagscan.engine.parse import TextParser
from tagscan.engine.textocr import _unrotate
from tagscan.profile import Profile, ListSpec


def test_normalize():
    assert normalize_number("28.462") == "28462"
    assert normalize_number(" 0104706 ") == "104706"
    assert normalize_number(104706.0) == "104706"
    assert normalize_number(None) == ""


def test_parser_keyword_and_brand():
    p = TextParser(Profile())
    r = p.parse("ACME\nASSET NO.\n123456")
    assert r.number == "123456" and r.level == "high" and r.brand and r.keyword
    r = p.parse("ASSET NO. 28.462\nMODEL X9")
    assert r.number == "28462" and r.level == "high"
    r = p.parse("random\n4711")
    assert r.number == "4711" and r.level == "low"
    assert p.parse("").number == ""


def test_parser_code_payload():
    p = TextParser(Profile())
    assert p.parse_code("Some Brand\nInventar-Nr.\n28019\nServicewagen") == "28019"
    assert p.parse_code("122145") == "122145"
    assert p.parse_code("12") == ""


def test_catalog_csv(tmp_path):
    f = tmp_path / "list.txt"
    f.write_text("Nr;Description\n28019;Servicewagen\n100016;Monitor\n", encoding="utf-8")
    cat = Catalog.load(f, ListSpec())
    assert len(cat) == 2 and "28.019" in cat and cat.get("100016").description == "Monitor"
    near = cat.nearest("100018", 1)
    assert near[0].key == "100016"


def test_unrotate_roundtrip():
    import numpy as np
    w0, h0 = 7, 4
    img = np.arange(w0 * h0).reshape(h0, w0)
    for k in range(4):
        rot = np.rot90(img, k)
        for y in range(rot.shape[0]):
            for x in range(rot.shape[1]):
                ox, oy = _unrotate(x, y, k, w0, h0)
                assert img[int(oy), int(ox)] == rot[y, x]


@pytest.mark.skipif(os.environ.get("TAGSCAN_NO_OCR") == "1", reason="OCR disabled")
def test_ocr_synthetic(tmp_path):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    import cv2
    from tagscan.app.main import synthetic_label
    from tagscan.engine import Scanner
    from tagscan.catalog import Item
    synthetic_label(tmp_path / "a.png", "123456")
    cat = Catalog(items={"123456": Item("123456", "123456", "thing")})
    sc = Scanner(Profile(), cat)
    img = cv2.imread(str(tmp_path / "a.png"))
    for rot in (0, 1, 2, 3):
        frame = __import__("numpy").ascontiguousarray(__import__("numpy").rot90(img, rot))
        sc.reset()
        obs, dec, _ = sc.process(frame, force_ocr=True)
        assert obs and obs[0].number == "123456", f"rotation {rot}"


def test_location_code_regex():
    p = TextParser(Profile())
    assert p.parse_location("LOC: Hall 3") == "Hall 3"
    assert p.parse_location("ort=Lab A\n") == "Lab A"
    assert p.parse_location("STANDORT:\n204") == "204"
    assert p.parse_location("Inventar-Nr.\n28019") == ""
    assert p.parse_location("") == ""
    assert p.parse_code("LOC: 4711") == "4711"       # parse_code alone would read it as a number ...


def test_location_code_wins_in_pipeline():
    """... which is why the scanner checks for a location payload before the number."""
    import cv2
    import numpy as np
    import zxingcpp
    from tagscan.engine import Scanner
    from tagscan.engine.codes import formats_from_names
    from tagscan.catalog import Item

    def qr(text):
        if hasattr(zxingcpp, "create_barcode"):                 # zxing-cpp >= 3
            img = np.array(zxingcpp.write_barcode_to_image(zxingcpp.create_barcode(text, zxingcpp.BarcodeFormat.QRCode), scale=8))
        else:
            img = np.array(zxingcpp.write_barcode(zxingcpp.BarcodeFormat.QRCode, text, width=360, height=360))
        img = cv2.copyMakeBorder(img, 40, 40, 40, 40, cv2.BORDER_CONSTANT, value=255)
        return cv2.cvtColor(img, cv2.COLOR_GRAY2BGR)

    cat = Catalog(items={"4711": Item("4711", "4711", "the item")})
    sc = Scanner(Profile(), cat)
    assert formats_from_names(Profile().code_formats) is not None

    obs, dec = sc.process_codes(qr("LOC: 4711"))
    assert dec.status == "location" and dec.observation.text == "4711" and obs[0].source == "LOCATION"
    obs, dec = sc.process_codes(qr("LOC: 4711"))           # the same code again is quiet for a while
    assert dec.status == "none" and obs and obs[0].source == "LOCATION"

    obs, dec = sc.process_codes(qr("4711"))                 # a plain number is still an item
    assert dec.status == "accept" and dec.observation.number == "4711" and dec.observation.in_list
    obs, dec = sc.process_codes(qr("Some Brand\nAsset No.\n5555\nthing"))
    assert dec.status == "review" and dec.observation.number == "5555"


def test_catalog_locations(tmp_path):
    f = tmp_path / "list.csv"
    f.write_text("No;Description;Location\n1001;Cart;Hall 10\n1002;Monitor;Hall 2\n1003;Mouse;Hall 2\n1004;x;\n",
                 encoding="utf-8")
    cat = Catalog.load(f, ListSpec(description_columns=[2], location_column=3))
    assert cat.get("1001").location == "Hall 10"
    assert cat.locations() == ["Hall 2", "Hall 10"]          # natural order, distinct, blanks dropped
    assert Catalog.load(f, ListSpec()).locations() == []


def test_profile_roundtrip_and_old_profiles(tmp_path):
    p = Profile()
    p.save(tmp_path / "p.json")
    q = Profile.load(tmp_path / "p.json")
    assert q.ocr_rotations == [0, 90] and q.code_formats == p.code_formats and q.list.location_column == 0
    # a 0.1 profile without the new keys still loads and gets the defaults
    old = {"name": "old", "list": {"sheet": "S", "header_row": 3}}
    r = Profile.from_dict(old)
    assert r.name == "old" and r.list.sheet == "S" and r.code_interval == 0.12 and r.camera_max_width == 1920


def test_scanner_apply_profile():
    from tagscan.engine import Scanner
    sc = Scanner(Profile())
    p2 = Profile(name="two", ocr_rotations=[0, 90, 180, 270], ocr_max_side=640, code_formats=[])
    p2.keyword_patterns = ["SERIAL"]
    sc.apply_profile(p2)
    assert sc.ocr.rotations == (0, 90, 180, 270) and sc.ocr.max_side == 640 and sc.formats is None
    assert sc.parser.parse("SERIAL 123456").level == "high" and sc.ocr.parser is sc.parser


def test_store_reads(tmp_path):
    import numpy as np
    from tagscan.store import Store
    st = Store(tmp_path)
    crop = np.full((40, 90, 3), 100, np.uint8)
    frame = np.full((120, 160, 3), 50, np.uint8)
    rid = st.add_read("typed", final_number="4711", source="MANUAL", ocr_number="4717", ocr_text="ASSET\n4717",
                      ocr_level="high", ocr_boxes=[[(1, 2), (3, 2), (3, 4), (1, 4)]], roi={"x": 1, "y": 2, "w": 3, "h": 4},
                      crop=crop, frame=frame, location="Hall 2")
    assert rid == 1
    row = st.db.execute("SELECT outcome, final_number, ocr_number, crop, photo, roi, ocr_boxes FROM reads").fetchone()
    assert row[0] == "typed" and row[1] == "4711" and row[2] == "4717"
    assert (tmp_path / "training" / "crops" / row[3]).exists() and (tmp_path / row[4]).exists()
    assert '"w": 3' in row[5] and row[6].startswith("[[[1.0, 2.0]")
    scan = st.add("4711", "thing", "Hall 2", "QR", True, image=frame)
    st.add_read("accepted", final_number="4711", source="QR", photo=scan.photo, frame=frame, scan_id=scan.id)
    row = st.db.execute("SELECT photo, scan_id FROM reads WHERE id=2").fetchone()
    assert row[0] == f"photos/{scan.photo}" and row[1] == scan.id       # no second copy of the frame
    assert st.read_counts() == {"typed": 1, "accepted": 1}


def test_operator_code_and_badge_in_pipeline():
    import cv2
    import numpy as np
    import zxingcpp
    from tagscan.engine import Scanner
    from tagscan.catalog import Item
    p = TextParser(Profile())
    assert p.parse_operator("OP: Anna Example") == "Anna Example"
    assert p.parse_operator("Mitarbeiter=4711") == "4711"
    assert p.parse_operator("LOC: Hall 3") == "" and p.parse_location("OP: Anna") == ""

    def qr(text):
        img = np.array(zxingcpp.write_barcode_to_image(zxingcpp.create_barcode(text, zxingcpp.BarcodeFormat.QRCode), scale=8))
        return cv2.cvtColor(cv2.copyMakeBorder(img, 40, 40, 40, 40, cv2.BORDER_CONSTANT, value=255), cv2.COLOR_GRAY2BGR)

    sc = Scanner(Profile(), Catalog(items={"4711": Item("4711", "4711", "x")}))
    obs, dec = sc.process_codes(qr("OP: Anna"))
    assert dec.status == "operator" and dec.observation.text == "Anna" and obs[0].source == "OPERATOR"
    obs, dec = sc.process_codes(qr("OP: Anna"))               # quiet while suppressed
    assert dec.status == "none" and obs[0].source == "OPERATOR"
    obs, dec = sc.process_codes(qr("4711"))                   # an item still resolves normally
    assert dec.status == "accept"


def test_store_photo_layout_operator_and_migration(tmp_path):
    import sqlite3
    import numpy as np
    from openpyxl import load_workbook
    from tagscan.catalog import Item
    from tagscan.store import Store, safe_name
    assert safe_name('Hall 2 / Bay "A": x?') == "Hall 2 _ Bay _A__ x"
    assert safe_name("   ") == "unknown" and safe_name("Büro 204") == "Büro 204" and safe_name("..") == "unknown"

    # a database written by 0.2 (no operator columns) is migrated on open
    (tmp_path / "old").mkdir()
    db = sqlite3.connect(tmp_path / "old" / "scans.sqlite")
    db.execute("CREATE TABLE scans(id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT, number TEXT, description TEXT,"
               " location TEXT, source TEXT, in_list INTEGER, photo TEXT, device TEXT, note TEXT DEFAULT '')")
    db.execute("INSERT INTO scans(ts,number,description,location,source,in_list,photo,device,note) "
               "VALUES('2026-10-16 10:00:00','1001','Cart','Hall 2','QR',1,'1001_20261016_100000.jpg','dev','')")
    db.commit()
    db.close()
    st = Store(tmp_path / "old")
    assert st.recent(1)[0].operator == "" and st.count() == 1

    st = Store(tmp_path / "new")
    frame = np.full((60, 80, 3), 90, np.uint8)
    a = st.add("1001", "Cart", "Hall 2", "QR", True, image=frame, operator="Anna")
    b = st.add("1001", "Cart", "Hall 2", "OCR", True, image=frame, operator="Anna")
    c = st.add("28.019", "x", "", "MANUAL", False, image=frame)
    import datetime as dt
    day = dt.datetime.now().strftime("%Y%m%d")
    assert a.photo == f"Hall 2/{day}_1001.jpg" and b.photo == f"Hall 2/{day}_1001_2.jpg"
    assert c.photo == f"no-location/{day}_28.019.jpg"
    for s in (a, b, c):
        assert (tmp_path / "new" / "photos" / s.photo).exists()
    assert st.recent(3)[0].operator == "" and st.recent(3)[2].operator == "Anna"
    st.delete(b.id)
    assert not (tmp_path / "new" / "photos" / b.photo).exists() and st.count() == 2
    st.add_read("accepted", final_number="1001", source="QR", photo=a.photo, scan_id=a.id, operator="Anna")
    assert st.db.execute("SELECT operator FROM reads").fetchone()[0] == "Anna"

    out = st.export_xlsx(tmp_path / "out.xlsx", Catalog(items={"1001": Item("1001", "1001", "Cart")}))
    wb = load_workbook(out)
    rows = list(wb["Scans"].iter_rows(values_only=True))
    assert rows[0][:5] == ("Number", "Description", "Kostenstelle", "Location", "Operator") and rows[1][4] == "Anna"
    status = list(wb["List status"].iter_rows(values_only=True))
    assert status[0][5] == "Operator" and status[1][3] == "found" and status[1][5] == "Anna"
    assert list(wb["Not in list"].iter_rows(values_only=True))[1][3] in ("", None)   # empty cell reads back as None
    assert "_Anna_" in st.default_export_name("Anna") and "__" not in st.default_export_name("")


def test_store_archive_starts_fresh(tmp_path):
    import numpy as np
    from tagscan.store import Store
    st = Store(tmp_path)
    frame = np.full((40, 60, 3), 90, np.uint8)
    a = st.add("1001", "Cart", "Hall 2", "QR", True, image=frame, operator="456cd")
    st.add_read("accepted", final_number="1001", source="QR", photo=a.photo, scan_id=a.id, crop=frame)
    st.export_xlsx(tmp_path / "TagScan_dev_x.xlsx")
    assert st.count() == 1 and (tmp_path / "photos" / a.photo).exists()
    dest = st.archive()
    assert dest.parent == tmp_path / "archive"
    assert st.count() == 0 and st.read_counts() == {} and not (tmp_path / "photos" / a.photo).exists()
    assert (dest / "photos" / a.photo).exists() and (dest / "scans.sqlite").exists()
    assert list((dest / "training" / "crops").iterdir()) and (dest / "TagScan_dev_x.xlsx").exists()
    assert not list(tmp_path.glob("*.xlsx"))
    b = st.add("1002", "Mouse", "Hall 2", "QR", True, image=frame)          # the fresh store works
    assert st.count() == 1 and (tmp_path / "photos" / b.photo).exists()


def test_describe_guess_from_label_text():
    p = TextParser(Profile())
    assert p.describe("ACME\nASSET NO.\n123456\nRolling cart 3 drawers", "123456") == "Rolling cart 3 drawers"
    assert p.describe("ACME\nASSET NO.\n123456", "123456") == ""
    assert p.describe("", "1") == "" and p.describe("12.03.2024\n123456\nMonitor 27", "123456") == "Monitor 27"
    assert len(p.describe("x" * 200, "")) == 60


def test_catalog_many_lists_kostenstelle(tmp_path):
    (tmp_path / "KST-100.csv").write_text("No;Description\n1001;Cart\n1002;Monitor\n", encoding="utf-8")
    (tmp_path / "KST-200.csv").write_text("No;Description\n2001;Laptop\n1002;Monitor again\n", encoding="utf-8")
    (tmp_path / "~$KST-200.csv").write_text("junk", encoding="utf-8")
    (tmp_path / "TagScan_dev_export.xlsx").write_bytes(b"not a list")
    files = Catalog.list_files(tmp_path)
    assert [f.name for f in files] == ["KST-100.csv", "KST-200.csv"]
    cat = Catalog.load_many(files, ListSpec(description_columns=[2]))
    assert len(cat) == 3 and cat.kst_names() == ["KST-100", "KST-200"] and cat.duplicates == 1
    assert cat.kst_of("1001") == "KST-100" and cat.kst_of("2001") == "KST-200" and cat.kst_of("1002") == "KST-100"
    assert cat.get("1002").description == "Monitor" and cat.kst_of("9999") == ""
    single = Catalog.load(tmp_path / "KST-100.csv", ListSpec(description_columns=[2]))
    assert single.kst_names() == ["KST-100"]


def test_store_export_per_kostenstelle(tmp_path):
    from openpyxl import load_workbook
    from tagscan.store import Store
    (tmp_path / "A.csv").write_text("No;Description\n1001;Cart\n1002;Monitor\n", encoding="utf-8")
    (tmp_path / "B.csv").write_text("No;Description\n2001;Laptop\n", encoding="utf-8")
    cat = Catalog.load_many(Catalog.list_files(tmp_path), ListSpec(description_columns=[2]))
    st = Store(tmp_path / "data")
    st.add("1001", "Cart", "Hall 2", "QR", True, operator="456cd", kst=cat.kst_of("1001"))
    st.add("2001", "Laptop", "Hall 3", "OCR", True, operator="456cd", kst=cat.kst_of("2001"))
    st.add("9999", "new thing", "Hall 3", "MANUAL", False, operator="456cd", kst="")
    assert st.recent(3)[2].kst == "A" and st.recent(3)[0].kst == ""
    files = st.export_all(tmp_path / "data" / "TagScan_dev_x.xlsx", cat)
    assert [f.name for f in files] == ["TagScan_dev_x.xlsx", "TagScan_dev_x_A.xlsx", "TagScan_dev_x_B.xlsx"]
    wb = load_workbook(files[0])
    scans = list(wb["Scans"].iter_rows(values_only=True))
    assert scans[0][2] == "Kostenstelle" and {r[2] or "" for r in scans[1:]} == {"A", "B", ""}
    status = list(wb["List status"].iter_rows(values_only=True))
    assert len(status) == 4 and {r[2] for r in status[1:]} == {"A", "B"} and [r[3] for r in status[1:]] == ["found", "missing", "found"]
    assert len(list(wb["Not in list"].iter_rows(values_only=True))) == 2
    wa = load_workbook(files[1])
    assert [r[0] for r in list(wa["Scans"].iter_rows(values_only=True))[1:]] == ["1001"]
    assert [r[3] for r in list(wa["List status"].iter_rows(values_only=True))[1:]] == ["found", "missing"]
    assert "Not in list" not in wa.sheetnames
    wb_b = load_workbook(files[2])
    assert [r[0] for r in list(wb_b["List status"].iter_rows(values_only=True))[1:]] == ["2001"]
    # a database from before 0.6 gets the column
    assert "kst" in [r[1] for r in st.db.execute("PRAGMA table_info(scans)")]
