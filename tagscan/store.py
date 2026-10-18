"""Session storage: one SQLite file plus a photos folder per device, and the Excel export.

Photos land in ``photos/<location>/<yyyymmdd>_<number>.jpg`` so the archive can be browsed by
room; a second photo of the same number on the same day there gets ``_2``, ``_3`` ..."""
from __future__ import annotations

import datetime as dt
import os
import re
import socket
import sqlite3
from dataclasses import dataclass
from pathlib import Path

from .catalog import Catalog


@dataclass
class Scan:
    id: int
    ts: str
    number: str
    description: str
    location: str
    source: str
    in_list: bool
    photo: str
    device: str
    note: str = ""
    operator: str = ""


def safe_name(text: str, fallback: str = "unknown", limit: int = 60) -> str:
    """A folder or file name Windows accepts: letters, digits, space, - _ . ( ) and accented
    letters survive, everything else becomes _ ; blank -> ``fallback``."""
    t = re.sub(r"[^\w \-.()]", "_", str(text or ""), flags=re.UNICODE)
    t = re.sub(r"\s+", " ", t).strip(" ._")
    return (t[:limit].rstrip(" ._") or fallback)


class Store:
    def __init__(self, folder: str | Path):
        self.folder = Path(folder)
        self.folder.mkdir(parents=True, exist_ok=True)
        (self.folder / "photos").mkdir(exist_ok=True)
        self.device = os.environ.get("TAGSCAN_DEVICE") or socket.gethostname()
        self.db = sqlite3.connect(self.folder / "scans.sqlite", check_same_thread=False)
        self.db.execute("""CREATE TABLE IF NOT EXISTS scans(
            id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT, number TEXT, description TEXT,
            location TEXT, source TEXT, in_list INTEGER, photo TEXT, device TEXT, note TEXT DEFAULT '',
            operator TEXT DEFAULT '')""")
        # every read the operator decided on, right or wrong: training data for a label detector
        # (frame + box + text boxes) and for OCR tuning (crop + the number that was actually there)
        self.db.execute("""CREATE TABLE IF NOT EXISTS reads(
            id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT, outcome TEXT, source TEXT,
            ocr_number TEXT, ocr_text TEXT, ocr_level TEXT, ocr_boxes TEXT, final_number TEXT,
            in_list INTEGER, roi TEXT, crop TEXT, photo TEXT, location TEXT, device TEXT, scan_id INTEGER,
            operator TEXT DEFAULT '')""")
        self._ensure_column("scans", "operator", "TEXT DEFAULT ''")      # databases from 0.1 / 0.2
        self._ensure_column("reads", "operator", "TEXT DEFAULT ''")
        self.db.commit()
        (self.folder / "training" / "crops").mkdir(parents=True, exist_ok=True)
        (self.folder / "training" / "frames").mkdir(parents=True, exist_ok=True)

    def _ensure_column(self, table: str, column: str, decl: str) -> None:
        cols = [r[1] for r in self.db.execute(f"PRAGMA table_info({table})")]
        if column not in cols:
            self.db.execute(f"ALTER TABLE {table} ADD COLUMN {column} {decl}")

    # ------------------------------------------------------------------ writes
    def photo_path(self, number: str, location: str, when: dt.datetime | None = None) -> str:
        """``<location>/<yyyymmdd>_<number>.jpg`` relative to the photos folder, unique."""
        when = when or dt.datetime.now()
        folder = safe_name(location, "no-location")
        base = f"{when.strftime('%Y%m%d')}_{safe_name(number, 'x')}"
        (self.folder / "photos" / folder).mkdir(parents=True, exist_ok=True)
        rel, n = f"{folder}/{base}.jpg", 1
        while (self.folder / "photos" / rel).exists():
            n += 1
            rel = f"{folder}/{base}_{n}.jpg"
        return rel

    def add(self, number: str, description: str, location: str, source: str, in_list: bool,
            image=None, note: str = "", operator: str = "") -> Scan:
        ts = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        photo = ""
        if image is not None:
            import cv2
            photo = self.photo_path(number, location)
            cv2.imwrite(str(self.folder / "photos" / photo), image, [int(cv2.IMWRITE_JPEG_QUALITY), 88])
        cur = self.db.execute("INSERT INTO scans(ts,number,description,location,source,in_list,photo,device,note,operator)"
                              " VALUES(?,?,?,?,?,?,?,?,?,?)",
                              (ts, number, description, location, source, int(in_list), photo, self.device, note, operator))
        self.db.commit()
        return Scan(cur.lastrowid, ts, number, description, location, source, in_list, photo, self.device, note, operator)

    def add_read(self, outcome: str, final_number: str = "", source: str = "", ocr_number: str = "",
                 ocr_text: str = "", ocr_level: str = "", ocr_boxes=None, roi: dict | None = None,
                 in_list: bool = False, crop=None, frame=None, photo: str = "", location: str = "",
                 scan_id: int | None = None, operator: str = "") -> int:
        """Record one decided read. ``outcome``: accepted | confirmed | corrected | typed | undone |
        skipped | back. ``crop`` (the box content) is saved under training/crops; the full
        ``frame`` under training/frames unless ``photo`` already names the scan's photo."""
        import json
        ts = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        crop_name, frame_name = "", ""
        if crop is not None and getattr(crop, "size", 0):
            import cv2
            crop_name = f"{stamp}.jpg"
            cv2.imwrite(str(self.folder / "training" / "crops" / crop_name), crop, [int(cv2.IMWRITE_JPEG_QUALITY), 90])
        if photo:
            frame_name = f"photos/{photo}"
        elif frame is not None and getattr(frame, "size", 0):
            import cv2
            frame_name = f"training/frames/{stamp}.jpg"
            cv2.imwrite(str(self.folder / frame_name), frame, [int(cv2.IMWRITE_JPEG_QUALITY), 85])
        boxes = json.dumps([[[round(float(x), 1), round(float(y), 1)] for x, y in q] for q in (ocr_boxes or [])])
        cur = self.db.execute(
            "INSERT INTO reads(ts,outcome,source,ocr_number,ocr_text,ocr_level,ocr_boxes,final_number,in_list,roi,"
            "crop,photo,location,device,scan_id,operator) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (ts, outcome, source, ocr_number, ocr_text or "", ocr_level, boxes, final_number, int(bool(in_list)),
             json.dumps(roi or {}), crop_name, frame_name, location, self.device, scan_id, operator))
        self.db.commit()
        return cur.lastrowid

    def archive(self) -> Path:
        """Start a new inventory: move the database, photos, training data and exports into
        archive/<date_time>/ and reopen empty. Settings, list and log stay."""
        import shutil
        self.db.close()
        dest = self.folder / "archive" / dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        dest.mkdir(parents=True, exist_ok=True)
        for name in ("scans.sqlite", "photos", "training"):
            src = self.folder / name
            if src.exists():
                shutil.move(str(src), str(dest / name))
        for f in self.folder.glob("*.xlsx"):
            shutil.move(str(f), str(dest / f.name))
        self.__init__(self.folder)
        return dest

    def read_counts(self) -> dict[str, int]:
        return {k: n for k, n in self.db.execute("SELECT outcome, COUNT(*) FROM reads GROUP BY outcome")}

    def delete(self, scan_id: int) -> None:
        row = self.db.execute("SELECT photo FROM scans WHERE id=?", (scan_id,)).fetchone()
        if row and row[0]:
            try:
                (self.folder / "photos" / row[0]).unlink()
            except OSError:
                pass
        self.db.execute("DELETE FROM scans WHERE id=?", (scan_id,))
        self.db.commit()

    def update(self, scan_id: int, **fields) -> None:
        if not fields:
            return
        cols = ", ".join(f"{k}=?" for k in fields)
        self.db.execute(f"UPDATE scans SET {cols} WHERE id=?", (*fields.values(), scan_id))
        self.db.commit()

    # ------------------------------------------------------------------ reads
    _COLS = "id,ts,number,description,location,source,in_list,photo,device,note,operator"

    def _rows(self, sql, args=()) -> list[Scan]:
        cur = self.db.execute(sql, args)
        return [Scan(r[0], r[1], r[2], r[3], r[4], r[5], bool(r[6]), r[7], r[8], r[9], r[10] or "") for r in cur.fetchall()]

    def recent(self, n: int = 12) -> list[Scan]:
        return self._rows(f"SELECT {self._COLS} FROM scans ORDER BY id DESC LIMIT ?", (n,))

    def all(self) -> list[Scan]:
        return self._rows(f"SELECT {self._COLS} FROM scans ORDER BY id")

    def count(self) -> int:
        return self.db.execute("SELECT COUNT(*) FROM scans").fetchone()[0]

    def seen(self, number: str) -> list[Scan]:
        return self._rows(f"SELECT {self._COLS} FROM scans WHERE number=? ORDER BY id", (number,))

    # ------------------------------------------------------------------ export
    def export_xlsx(self, path: str | Path, catalog: Catalog | None = None) -> Path:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "Scans"
        head = ["Number", "Description", "Location", "Operator", "Time", "Source", "In list", "Photo", "Device", "Note"]
        ws.append(head)
        for s in self.all():
            ws.append([s.number, s.description, s.location, s.operator, s.ts, s.source, "yes" if s.in_list else "no",
                       s.photo, s.device, s.note])
        _style(ws, head, Font, PatternFill, get_column_letter)

        if catalog and len(catalog):
            found = {s.number for s in self.all()}
            ws2 = wb.create_sheet("List status")
            head2 = ["Number", "Description", "Status", "Location", "Operator", "Time", "Device"]
            ws2.append(head2)
            last = {}
            for s in self.all():
                last[s.number] = s
            for it in catalog.items.values():
                s = last.get(it.key)
                ws2.append([it.number, it.description, "found" if s else "missing",
                            s.location if s else "", s.operator if s else "", s.ts if s else "", s.device if s else ""])
            _style(ws2, head2, Font, PatternFill, get_column_letter)
            ws3 = wb.create_sheet("Not in list")
            head3 = ["Number", "Description", "Location", "Operator", "Time", "Source", "Photo", "Device", "Note"]
            ws3.append(head3)
            for s in self.all():
                if not s.in_list:
                    ws3.append([s.number, s.description, s.location, s.operator, s.ts, s.source, s.photo, s.device, s.note])
            _style(ws3, head3, Font, PatternFill, get_column_letter)
        path = Path(path)
        wb.save(path)
        return path

    def default_export_name(self, operator: str = "") -> str:
        who = f"_{safe_name(operator, '')}" if safe_name(operator, "") else ""
        return f"TagScan_{self.device}{who}_{dt.datetime.now().strftime('%Y-%m-%d_%H%M')}.xlsx"


def _style(ws, head, Font, PatternFill, get_column_letter):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    widths = {"Number": 12, "Description": 48, "Location": 18, "Operator": 16, "Time": 20, "Source": 9, "In list": 8,
              "Photo": 36, "Device": 16, "Note": 30, "Status": 9}
    for i, h in enumerate(head, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 14)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
