"""The list of known items (xlsx, csv or txt) and lookups against it."""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

from .profile import ListSpec


def normalize_number(raw) -> str:
    """'28.462' -> '28462', ' 0104706 ' -> '104706', 104706.0 -> '104706'."""
    if raw is None:
        return ""
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)
    s = re.sub(r"[.\s]", "", str(raw).strip())
    return s.lstrip("0") or ("0" if s else "")


@dataclass
class Item:
    number: str                   # as written in the list
    key: str                      # normalized
    description: str = ""
    extra: list = field(default_factory=list)
    location: str = ""            # from the profile's location_column, if any
    kst: str = ""                 # Kostenstelle: the name of the list file the item came from


@dataclass
class Catalog:
    items: dict[str, Item] = field(default_factory=dict)
    headers: list[str] = field(default_factory=list)
    source: str = ""
    duplicates: int = 0           # numbers that appeared in more than one list (the first one wins)

    def __len__(self) -> int:
        return len(self.items)

    def __contains__(self, number) -> bool:
        return normalize_number(number) in self.items

    def get(self, number) -> Item | None:
        return self.items.get(normalize_number(number))

    def kst_names(self) -> list[str]:
        """The Kostenstellen (list files) the items came from, in loading order."""
        return list(dict.fromkeys(it.kst for it in self.items.values() if it.kst))

    def kst_of(self, number) -> str:
        it = self.get(number)
        return it.kst if it else ""

    def locations(self) -> list[str]:
        """Distinct locations named in the list (profile ``list.location_column``), naturally sorted."""
        seen = {it.location: None for it in self.items.values() if it.location}
        return sorted(seen, key=_natural_key)

    def nearest(self, number, k: int = 5) -> list[Item]:
        key = normalize_number(number)
        if not key:
            return []
        try:
            target = int(key)
        except ValueError:
            target = None
        scored = []
        for it in self.items.values():
            if target is not None and it.key.isdigit():
                d = abs(target - int(it.key))
                # a one-digit OCR slip (6<->8, 1<->7) should rank before a numerically closer number
                if len(it.key) == len(key):
                    d = min(d, 10 * sum(a != b for a, b in zip(it.key, key)))
            else:
                d = _edit_distance(key, it.key) * 100
            scored.append((d, it))
        scored.sort(key=lambda x: x[0])
        return [it for _, it in scored[:k]]

    # ------------------------------------------------------------------ loading
    @classmethod
    def load(cls, path: str | Path, spec: ListSpec | None = None, kst: str | None = None) -> "Catalog":
        """One list. ``kst`` (Kostenstelle) defaults to the file name without extension."""
        spec = spec or ListSpec()
        path = Path(path)
        cat = cls._load_xlsx(path, spec) if path.suffix.lower() in (".xlsx", ".xlsm") else cls._load_csv(path, spec)
        name = path.stem if kst is None else kst
        for it in cat.items.values():
            it.kst = name
        return cat

    @classmethod
    def load_many(cls, paths, spec: ListSpec | None = None) -> "Catalog":
        """Several lists, one per Kostenstelle (the file name). A number in two lists keeps the
        first one and is counted in ``duplicates``."""
        paths = [Path(p) for p in paths]
        cat = cls(source="; ".join(str(p) for p in paths))
        for path in paths:
            part = cls.load(path, spec)
            if not cat.headers:
                cat.headers = part.headers
            for key, it in part.items.items():
                if key in cat.items:
                    cat.duplicates += 1
                    continue
                cat.items[key] = it
        return cat

    @staticmethod
    def list_files(folder: str | Path) -> list[Path]:
        """The list files in a folder: xlsx/xlsm/csv/txt, not temporary (~$) or exported ones."""
        folder = Path(folder)
        out = []
        for p in sorted(folder.iterdir()):
            if p.is_file() and p.suffix.lower() in (".xlsx", ".xlsm", ".csv", ".txt") \
                    and not p.name.startswith("~$") and not p.name.startswith("TagScan_"):
                out.append(p)
        return out

    @staticmethod
    def sheet_names(path: str | Path) -> list[str]:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()

    @classmethod
    def _load_xlsx(cls, path: Path, spec: ListSpec) -> "Catalog":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[spec.sheet] if spec.sheet and spec.sheet in wb.sheetnames else wb[wb.sheetnames[0]]   # sheet missing: first
            cat = cls(source=str(path))
            ncol = max([spec.number_column, *spec.description_columns])
            for i, row in enumerate(ws.iter_rows(min_row=1, max_col=ncol, values_only=True), start=1):
                if i == spec.header_row:
                    cat.headers = [str(v) if v is not None else f"Col {j+1}" for j, v in enumerate(row)]
                    continue
                if i <= spec.header_row:
                    continue
                cat._add_row(list(row), spec)
            return cat
        finally:
            wb.close()

    @classmethod
    def _load_csv(cls, path: Path, spec: ListSpec) -> "Catalog":
        text = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = path.read_text(encoding=enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError(f"cannot decode {path}")
        delim = spec.delimiter
        if not delim:
            head = text[:2000]
            delim = max([";", ",", "\t", "|"], key=head.count)
        cat = cls(source=str(path))
        for i, row in enumerate(csv.reader(text.splitlines(), delimiter=delim), start=1):
            if i == spec.header_row:
                cat.headers = row
                continue
            if i <= spec.header_row or not row:
                continue
            cat._add_row(row, spec)
        return cat

    def _add_row(self, row: list, spec: ListSpec) -> None:
        def cell(c):
            return row[c - 1] if 0 < c <= len(row) else None
        raw = cell(spec.number_column)
        key = normalize_number(raw)
        if not key or not re.fullmatch(r"\d+", key):
            return
        desc_parts = [str(cell(c)).strip() for c in spec.description_columns if cell(c) not in (None, "")]
        number = str(int(raw)) if isinstance(raw, float) and raw.is_integer() else str(raw).strip()
        loc = cell(spec.location_column) if spec.location_column else None
        self.items[key] = Item(number=number, key=key,
                               description=" | ".join(desc_parts),
                               extra=[cell(c) for c in spec.description_columns],
                               location="" if loc in (None, "") else str(loc).strip())


def _natural_key(s: str):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", s)]


def _edit_distance(a: str, b: str) -> int:
    if len(a) < len(b):
        a, b = b, a
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a):
        cur = [i + 1]
        for j, cb in enumerate(b):
            cur.append(min(cur[j] + 1, prev[j + 1] + 1, prev[j] + (ca != cb)))
        prev = cur
    return prev[-1]
